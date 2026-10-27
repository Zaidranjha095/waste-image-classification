"""Write the report, slides, voiceover script, LinkedIn draft, checklist, and README.

Run after train.py:
    python build_docs.py
"""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
INK = RGBColor(0x1A, 0x24, 0x21)
GREEN = RGBColor(0x1F, 0x4D, 0x3A)
AMBER = RGBColor(0xC4, 0x7B, 0x2B)
CREAM = RGBColor(0xF3, 0xEF, 0xE6)
PAPER = RGBColor(0xFF, 0xFD, 0xF8)


def pct(value: float) -> str:
    return f"{100 * float(value):.2f}%"


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def confusion_pairs(metrics: dict) -> list[tuple[int, str, str]]:
    classes = metrics["class_order"]
    matrix = metrics["test"]["confusion_matrix"]
    pairs = []
    for row, true_name in enumerate(classes):
        for column, pred_name in enumerate(classes):
            count = int(matrix[row][column])
            if row != column and count:
                pairs.append((count, true_name, pred_name))
    pairs.sort(key=lambda item: item[0], reverse=True)
    return pairs


def confusion_sentence(metrics: dict) -> str:
    pairs = confusion_pairs(metrics)
    if not pairs:
        return "Every test image was classified correctly."
    bits = [f"{count} {true_name} images predicted as {pred_name}" for count, true_name, pred_name in pairs[:3]]
    return "The largest confusions were " + "; ".join(bits) + "."


def phase_epochs(metrics: dict, phase: str) -> int:
    return sum(1 for row in metrics["history"] if row["phase"] == phase)


def student_intro(config: dict) -> str:
    name = config["studentName"].strip()
    roll = config["rollNumber"].strip()
    if name == "Student Name" or roll == "Roll Number":
        return "Start by saying your name, your roll number, and that this is your Deep Learning semester project."
    return f"My name is {name}, roll number {roll}. This is my Deep Learning semester project."


def display_url(config: dict, key: str, label: str) -> str:
    value = str(config.get(key, "")).strip()
    return value if value else f"[{label} — add this before submission]"


def write_report(config: dict, metrics: dict) -> None:
    classes = metrics["class_order"]
    counts = metrics["counts"]
    training = metrics["training"]
    test = metrics["test"]
    weights = training["class_weights"]
    head_epochs = phase_epochs(metrics, "head")
    fine_epochs = phase_epochs(metrics, "finetune")
    device = training["gpu"] or "CPU"
    count_bits = ", ".join(f"{name} {counts['all'][name]}" for name in classes)
    weight_bits = ", ".join(f"{name} {weights[name]:.3f}" for name in classes)
    header = "| Class | Precision | Recall | F1 | Support |"
    divider = "| --- | --- | --- | --- | --- |"
    table_rows = [
        f"| {name} | {pct(test['per_class'][name]['precision'])} | {pct(test['per_class'][name]['recall'])} | {pct(test['per_class'][name]['f1'])} | {test['per_class'][name]['support']} |"
        for name in classes
    ]
    lines = [
        f"# {config['title']}",
        "",
        f"**Student:** {config['studentName']}",
        "",
        f"**Roll number:** {config['rollNumber']}",
        "",
        f"**Course:** {config['course']}",
        "",
        "## Problem statement",
        "",
        "Household and campus waste is easier to recycle when paper, plastic, glass, metal, and cardboard are separated before disposal. Doing that by hand does not scale, and a photograph is often the only signal available at a bin or a conveyor.",
        "",
        "This project builds a six-class image classifier for that decision. The classes are cardboard, glass, metal, paper, plastic, and trash. The trained network is saved as a PyTorch checkpoint and as an ONNX file so the same weights can be downloaded for evaluation and run in a browser demo.",
        "",
        "The system is an educational demonstration. It is not a controller for hazardous-waste handling, and it was not trained on medical, chemical, or industrial-safety categories.",
        "",
        "## Objectives",
        "",
        "- Train a convolutional classifier on a public waste-image dataset.",
        "- Keep a held-out test set out of model selection and report accuracy, per-class F1, and a confusion matrix.",
        "- Publish the architecture, checkpoint, and a one-command prediction script.",
        "- Deploy a web page that shows the dataset, the code layout, live predictions, this report, and the presentation.",
        "",
        "## Dataset",
        "",
        "The images come from TrashNet, collected by Gary Thung and Mindy Yang for Stanford CS 229 and published at [github.com/garythung/trashnet](https://github.com/garythung/trashnet). Each photo shows one object on a plain posterboard under sunlight or room light. The phones used for the original collection were an iPhone 7 Plus, an iPhone 5S, and an iPhone SE. This project uses the resized JPEG release (`dataset-resized.zip`), whose images are 512 by 384 pixels.",
        "",
        f"The release contains **{counts['total']}** images: {count_bits}. Trash is the smallest class, with {counts['all']['trash']} images, so a classifier that ignores it can still post a high overall accuracy. Training therefore uses inverse-frequency class weights, and the report includes per-class scores rather than accuracy alone.",
        "",
        f"The split is stratified with seed {metrics['seed']}: **{counts['train']['total']}** train, **{counts['val']['total']}** validation, and **{counts['test']['total']}** test, about 70% / 15% / 15%. Early stopping watches validation accuracy. The test set is used once, after the checkpoint is chosen.",
        "",
        "The full image archive is not committed. `download_dataset.py` fetches it, and `samples/` keeps one correctly classified test image per class for the website and the slides.",
        "",
        "![cardboard](../samples/cardboard.jpg)",
        "",
        "![glass](../samples/glass.jpg)",
        "",
        "![metal](../samples/metal.jpg)",
        "",
        "![paper](../samples/paper.jpg)",
        "",
        "![plastic](../samples/plastic.jpg)",
        "",
        "![trash](../samples/trash.jpg)",
        "",
        "## Architecture",
        "",
        "The network is `torchvision.models.mobilenet_v2`, initialized from ImageNet-1K weights (`IMAGENET1K_V1`). MobileNetV2 is a stack of inverted residual blocks. A block expands the channels with a 1×1 convolution, filters them with a depthwise 3×3 convolution, and projects them back with a linear 1×1 convolution. Stride-1 blocks add a residual connection when the depth does not change. That design keeps the model small enough to store in Git and to run in the browser.",
        "",
        f"The ImageNet classifier is replaced by dropout (probability 0.2, the torchvision default) and a linear layer from 1,280 features to 6 classes. The full network has **{training['parameter_count']:,}** parameters. Inputs are RGB images. Evaluation resizes the shorter side to 256, takes a center crop of 224, and normalizes with the ImageNet mean (0.485, 0.456, 0.406) and standard deviation (0.229, 0.224, 0.225). The exported model returns logits. Probabilities are a softmax over those six logits.",
        "",
        "`model.py` defines `build_model()`. Loading `weights/model.pth` into that module reproduces the trained network. The checkpoint dictionary stores `state_dict`, `classes`, `input_size`, `mean`, `std`, and `arch`.",
        "",
        "## Training",
        "",
        f"Training ran on **{device}** with batch size {training['batch_size']}. The loss is cross-entropy with inverse-frequency weights computed on the training split only: {weight_bits}. The optimizer is Adam with weight decay {training['weight_decay']}. Each phase also uses ReduceLROnPlateau on validation accuracy, with factor 0.5 and patience 2.",
        "",
        "Training images are augmented with a random resized crop (scale 0.7 to 1.0), horizontal flip, rotation up to 15 degrees, and color jitter of 0.2 brightness, 0.2 contrast, 0.2 saturation, and 0.05 hue. Validation and test images use the deterministic resize-and-center-crop pipeline.",
        "",
        f"Phase one freezes every feature block and trains the new classifier for up to 8 epochs at learning rate 1e-3, with early stopping patience 3. It ran for **{head_epochs}** epochs. The best head checkpoint is reloaded before fine-tuning.",
        "",
        f"Phase two unfreezes MobileNetV2 feature blocks from index 14 onward (the later inverted residuals and the final convolution) and trains them at learning rate 1e-4 while the classifier stays at 1e-3. The cap is 20 epochs and the patience is 5. It ran for **{fine_epochs}** epochs. The saved weights are the epoch with the highest validation accuracy, **{pct(training['best_val_accuracy'])}**, across both phases. In total, **{training['epochs_ran']}** epochs were executed.",
        "",
        "![Training curves](../outputs/training_curves.png)",
        "",
        "## Results",
        "",
        f"On the held-out test set the model reaches **{pct(test['accuracy'])}** accuracy, **{pct(test['macro_f1'])}** macro F1, and **{pct(test['weighted_f1'])}** weighted F1. Macro F1 treats the six classes equally, so the small trash class cannot hide inside the overall accuracy.",
        "",
        header,
        divider,
        *table_rows,
        "",
        "![Test confusion matrix](../outputs/confusion_matrix.png)",
        "",
        confusion_sentence(metrics),
        "",
        "The original TrashNet course project reported roughly 75% test accuracy with a smaller convolutional network trained from scratch, on the authors' own split. That figure is context, not a paired comparison: this project uses ImageNet initialization and a different stratified split.",
        "",
        "## Web deployment and repository",
        "",
        "The static site in `web/` states the problem, the dataset, the code layout, and these results. Uploading a photo, or clicking a sample, runs `model.onnx` with ONNX Runtime Web. The page also links to this report, the slide deck, and the weight files.",
        "",
        "Canonical files for an automated check:",
        "",
        "- `weights/model.pth` — PyTorch checkpoint",
        "- `weights/model.onnx` — browser copy of the same network",
        "- `weights/config.json` — class order, input size, and normalization",
        "- `weights/class_names.json` — class order alone",
        "- `model_manifest.json` — paths and the prediction command",
        "- `model.py` — `build_model()`",
        "- `predict.py` — prints JSON probabilities for one image",
        "- `outputs/metrics.json` — the numbers in this report",
        "",
        "Command-line prediction:",
        "",
        "```",
        "python predict.py --image samples/paper.jpg",
        "```",
        "",
        "Loading the checkpoint in Python:",
        "",
        "```python",
        "import torch",
        "from model import build_model",
        "",
        'checkpoint = torch.load("weights/model.pth", map_location="cpu", weights_only=False)',
        'model = build_model(num_classes=len(checkpoint["classes"]), pretrained=False)',
        'model.load_state_dict(checkpoint["state_dict"])',
        "model.eval()",
        "```",
        "",
        "Class order in the checkpoint and in the ONNX softmax is cardboard, glass, metal, paper, plastic, trash.",
        "",
        "## Limitations",
        "",
        "- TrashNet objects sit on a clean posterboard. A photo of a bin, a street, or several overlapping items is outside that distribution, so the demo can be confidently wrong.",
        "- Glass, plastic, and metal are often transparent or reflective, and the confusion matrix is where those mistakes show up.",
        "- The trash class has 137 images in the whole dataset and 20 in the test split, so its F1 is a noisy estimate.",
        "- Inverse-frequency weights push the model to notice trash, but they do not create new photographs.",
        "- Nothing in the label set covers batteries, chemicals, or sharp medical waste. The page should not be used as a safety instruction.",
        "",
        "## Reproducibility",
        "",
        "```",
        "python -m venv .venv",
        ".venv\\Scripts\\python -m pip install torch torchvision --index-url https://download.pytorch.org/whl/cu128",
        ".venv\\Scripts\\python -m pip install -r requirements.txt",
        ".venv\\Scripts\\python download_dataset.py",
        ".venv\\Scripts\\python train.py",
        ".venv\\Scripts\\python build_docs.py",
        "node prepare_site.js",
        "```",
        "",
        "CPU training also finishes, because the dataset is small. The CUDA 12.8 wheel is what this machine used for an RTX 5070 (Blackwell, `sm_120`). `train.py` falls back to CPU if a CUDA kernel cannot run.",
        "",
        "## References",
        "",
        "1. Gary Thung and Mindy Yang. TrashNet. Stanford CS 229, 2016. https://github.com/garythung/trashnet",
        "",
        "2. Mark Sandler, Andrew Howard, Menglong Zhu, Andrey Zhmoginov, and Liang-Chieh Chen. MobileNetV2: Inverted Residuals and Linear Bottlenecks. CVPR, 2018.",
        "",
        "3. Torchvision MobileNetV2 and ImageNet-1K weights, `MobileNet_V2_Weights.IMAGENET1K_V1`.",
        "",
    ]
    report_dir = ROOT / "report"
    report_dir.mkdir(parents=True, exist_ok=True)
    (report_dir / "PROJECT_REPORT.md").write_text("\n".join(lines), encoding="utf-8")


def add_rect(slide, left, top, width, height, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def style_run(run, size: int, bold: bool, color: RGBColor, font: str = "Calibri") -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text(slide, text, left, top, width, height, size=20, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    style_run(run, size, bold, color)
    return box


def add_bullets(slide, bullets: list[str], left, top, width, height, size=22) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = "•  " + bullet
        style_run(run, size, False, INK)
        paragraph.space_after = Pt(10)


def set_notes(slide, notes: str) -> None:
    slide.notes_slide.notes_text_frame.text = notes


def blank_slide(prs: Presentation, title: str, notes: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = CREAM
    add_rect(slide, 0, 0, prs.slide_width, Inches(0.16), GREEN)
    add_text(slide, title, Inches(0.5), Inches(0.35), Inches(12.3), Inches(0.7), size=32, bold=True, color=GREEN)
    set_notes(slide, notes)
    return slide


def set_cell(cell, text: str, size=14, bold=False, fill: RGBColor | None = None, font_color: RGBColor = INK) -> None:
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    frame = cell.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    style_run(run, size, bold, font_color)


def write_slides(config: dict, metrics: dict, notes_by_title: dict[str, str]) -> None:
    classes = metrics["class_order"]
    test = metrics["test"]
    training = metrics["training"]
    counts = metrics["counts"]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = config["title"]
    prs.core_properties.author = config["studentName"]

    title = prs.slides.add_slide(prs.slide_layouts[6])
    fill = title.background.fill
    fill.solid()
    fill.fore_color.rgb = GREEN
    add_text(title, "DEEP LEARNING", Inches(0.7), Inches(1.3), Inches(11), Inches(0.4), size=16, bold=True, color=AMBER)
    add_text(title, "Waste Image Classification", Inches(0.7), Inches(1.8), Inches(12), Inches(1.3), size=48, bold=True, color=PAPER)
    add_text(
        title,
        "MobileNetV2 transfer learning on TrashNet",
        Inches(0.7),
        Inches(3.3),
        Inches(11),
        Inches(0.6),
        size=24,
        color=CREAM,
    )
    add_text(
        title,
        f"{config['studentName']}  ·  {config['rollNumber']}",
        Inches(0.7),
        Inches(5.6),
        Inches(11),
        Inches(0.5),
        size=20,
        color=PAPER,
    )
    set_notes(title, notes_by_title["Title"])

    slide = blank_slide(prs, "The decision this model makes", notes_by_title["Problem"])
    add_bullets(
        slide,
        [
            "Recycling works when cardboard, paper, plastic, glass, and metal are separated.",
            "The model looks at one photograph and chooses one of six classes.",
            "The sixth class, trash, is residual material that does not fit the other five.",
            "The result is a semester demo with downloadable weights, not a plant controller.",
        ],
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(5),
    )

    slide = blank_slide(prs, "What the project delivers", notes_by_title["Objectives"])
    add_bullets(
        slide,
        [
            "A public secondary dataset, TrashNet, with a fixed stratified split.",
            "A fine-tuned MobileNetV2 checkpoint and a matching ONNX file.",
            "predict.py, which prints class probabilities for one image.",
            "A Vercel page with the demo, this deck, and the written report.",
        ],
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(5),
    )

    slide = blank_slide(prs, "TrashNet", notes_by_title["Dataset"])
    add_bullets(
        slide,
        [
            "Gary Thung and Mindy Yang, Stanford CS 229. Resized JPEGs, 512×384.",
            f"{counts['total']} images: " + ", ".join(f"{name} {counts['all'][name]}" for name in classes) + ".",
            f"Stratified split, seed {metrics['seed']}: {counts['train']['total']} train, {counts['val']['total']} validation, {counts['test']['total']} test.",
            "Objects sit on a plain posterboard. That is both why the task is learnable and why wild photos are harder.",
        ],
        Inches(0.6),
        Inches(1.4),
        Inches(12),
        Inches(5.4),
        size=20,
    )

    slide = blank_slide(prs, "One test image from each class", notes_by_title["Examples"])
    for index, name in enumerate(classes):
        left = Inches(0.4 + index * 2.15)
        picture = str(ROOT / "samples" / f"{name}.jpg")
        slide.shapes.add_picture(picture, left, Inches(1.6), width=Inches(2.0))
        add_text(slide, name, left, Inches(4.5), Inches(2.0), Inches(0.4), size=16, bold=True, color=GREEN, align=PP_ALIGN.CENTER)

    slide = blank_slide(prs, "MobileNetV2 with a six-way head", notes_by_title["Architecture"])
    add_bullets(
        slide,
        [
            "ImageNet-1K initialization, then a new linear layer with 6 outputs.",
            "Inverted residual blocks: expand, depthwise 3×3, linear project.",
            f"{training['parameter_count']:,} parameters. Small enough for Git and for the browser.",
            "Input 224×224 after a shorter-side resize to 256 and a center crop. Output is logits.",
        ],
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(5),
        size=22,
    )

    slide = blank_slide(prs, "Two-phase fine-tuning", notes_by_title["Training"])
    add_bullets(
        slide,
        [
            f"Device: {training['gpu'] or 'CPU'}. Batch size {training['batch_size']}. Adam, weight decay {training['weight_decay']}.",
            "Loss: class-weighted cross-entropy. Trash receives the largest weight.",
            f"Head only, then feature blocks from index 14. {training['epochs_ran']} epochs ran.",
            f"Best validation accuracy {pct(training['best_val_accuracy'])}. Test labels were not used to pick the epoch.",
        ],
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(5),
        size=22,
    )

    slide = blank_slide(prs, "Held-out test results", notes_by_title["Results"])
    add_text(
        slide,
        f"{pct(test['accuracy'])} accuracy     {pct(test['macro_f1'])} macro F1",
        Inches(0.6),
        Inches(1.25),
        Inches(12),
        Inches(0.6),
        size=28,
        bold=True,
        color=GREEN,
    )
    rows, cols = 1 + len(classes), 5
    table_shape = slide.shapes.add_table(rows, cols, Inches(0.6), Inches(2.15), Inches(12), Inches(4.6))
    table = table_shape.table
    headers = ["Class", "Precision", "Recall", "F1", "Support"]
    for column, header in enumerate(headers):
        set_cell(table.cell(0, column), header, size=14, bold=True, fill=GREEN, font_color=PAPER)
    for row, name in enumerate(classes, start=1):
        stats = test["per_class"][name]
        values = [name, pct(stats["precision"]), pct(stats["recall"]), pct(stats["f1"]), str(stats["support"])]
        for column, value in enumerate(values):
            set_cell(table.cell(row, column), value, size=15, bold=column == 0)

    slide = blank_slide(prs, "Learning curves", notes_by_title["Curves"])
    slide.shapes.add_picture(str(ROOT / "outputs" / "training_curves.png"), Inches(1.4), Inches(1.4), width=Inches(10.4))

    slide = blank_slide(prs, "Where the errors go", notes_by_title["Confusion"])
    slide.shapes.add_picture(str(ROOT / "outputs" / "confusion_matrix.png"), Inches(0.5), Inches(1.35), width=Inches(6.2))
    add_text(slide, confusion_sentence(metrics), Inches(6.9), Inches(2.0), Inches(5.8), Inches(3.5), size=20, color=INK)

    slide = blank_slide(prs, "What an evaluator can download", notes_by_title["Demo"])
    add_bullets(
        slide,
        [
            "weights/model.pth — load state_dict with build_model() from model.py.",
            "weights/model.onnx — the file the Vercel page runs in the browser.",
            "predict.py --image <path> prints JSON with the class and six probabilities.",
            "Class order is cardboard, glass, metal, paper, plastic, trash.",
        ],
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(5),
        size=22,
    )

    slide = blank_slide(prs, "What this model should not be asked to do", notes_by_title["Limits"])
    add_bullets(
        slide,
        [
            "TrashNet backgrounds are plain. A real bin is a different problem.",
            "Glass, plastic, and metal reflect light, so they trade errors.",
            f"Trash has only {counts['all']['trash']} images in the whole dataset.",
            "Do not use the demo to decide whether waste is hazardous.",
        ],
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(5),
    )

    slide = blank_slide(prs, "Links and closing", notes_by_title["Close"])
    add_bullets(
        slide,
        [
            "Dataset: github.com/garythung/trashnet",
            f"GitHub: {display_url(config, 'githubUrl', 'GitHub URL')}",
            f"Vercel: {display_url(config, 'vercelUrl', 'Vercel URL')}",
            f"YouTube: {display_url(config, 'youtubeUrl', 'YouTube URL')}",
            f"LinkedIn: {display_url(config, 'linkedinUrl', 'LinkedIn URL')}",
        ],
        Inches(0.6),
        Inches(1.45),
        Inches(12),
        Inches(5),
        size=20,
    )

    out = ROOT / "presentation" / "presentation.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)


def write_voiceover(notes_by_title: dict[str, str]) -> None:
    lines = [
        "# Voiceover script",
        "",
        "Record the slide deck in order. Read one section while that slide is on screen, then advance. Speak a little slower than conversation. The same text is stored in the PowerPoint speaker notes.",
        "",
        "Export the recording and upload it to YouTube as Public or Unlisted. Put the URL in `project.config.json` under `youtubeUrl`, run `node prepare_site.js`, and redeploy.",
        "",
    ]
    for title, notes in notes_by_title.items():
        lines.extend([f"## {title}", "", notes, ""])
    path = ROOT / "presentation" / "VOICEOVER_SCRIPT.md"
    path.write_text("\n".join(lines), encoding="utf-8")


def write_submission(config: dict, metrics: dict) -> None:
    test = metrics["test"]
    github = display_url(config, "githubUrl", "GitHub repository URL")
    vercel = display_url(config, "vercelUrl", "Vercel URL")
    youtube = display_url(config, "youtubeUrl", "YouTube URL")
    linkedin = display_url(config, "linkedinUrl", "LinkedIn URL")
    folder = ROOT / "submission"
    folder.mkdir(parents=True, exist_ok=True)
    post = f"""# LinkedIn post

Copy the text below. Replace the bracketed URLs if they are still placeholders, then publish the post and paste its URL into `project.config.json` (`linkedinUrl`).

---

I trained a waste-image classifier for my Deep Learning semester project.

The model is MobileNetV2, fine-tuned on TrashNet: cardboard, glass, metal, paper, plastic, and trash. On a held-out test split it reached {pct(test['accuracy'])} accuracy and {pct(test['macro_f1'])} macro F1. The checkpoint is small enough to publish, and the project page runs the same network in the browser.

- Live demo and report: {vercel}
- Code, dataset link, and weights: {github}
- Walkthrough: {youtube}

Weights are at `weights/model.pth`. TrashNet photographs use a plain background, so the demo is a course project, not a recycling-plant system.

#DeepLearning #ComputerVision #MachineLearning
"""
    (folder / "LINKEDIN_POST.md").write_text(post, encoding="utf-8")
    checklist = f"""# Submission checklist

Measured test accuracy: {pct(test['accuracy'])}. Macro F1: {pct(test['macro_f1'])}.

- [ ] Put your name and roll number in `project.config.json`.
- [ ] Re-run `python build_docs.py` and `node prepare_site.js` so the report and slides use that name.
- [ ] Create a public GitHub repository and push this project. Do not enable Git LFS for the weights.
- [ ] Confirm these files open from the repository: `weights/model.pth`, `weights/model.onnx`, `weights/config.json`, `model.py`, `predict.py`, `model_manifest.json`.
- [ ] Put the repository URL in `project.config.json` as `githubUrl`.
- [ ] Deploy the site to Vercel from this folder (`npx vercel`). The build command is `node prepare_site.js`.
- [ ] Open the Vercel URL and check the demo, `report.html`, `downloads/model.pth`, and `downloads/presentation.pptx`.
- [ ] Put the Vercel URL in `project.config.json` as `vercelUrl`.
- [ ] Record the voiceover from `presentation/VOICEOVER_SCRIPT.md` while showing `presentation/presentation.pptx`.
- [ ] Upload the video to YouTube as Public or Unlisted. Store the URL as `youtubeUrl`.
- [ ] Publish `submission/LINKEDIN_POST.md` and store the post URL as `linkedinUrl`.
- [ ] Run `node prepare_site.js` again and redeploy so the live page contains the GitHub, YouTube, and LinkedIn links.
- [ ] Submit the instructor form with the Vercel URL, GitHub repository, YouTube video, and LinkedIn post.

Current placeholders:

- GitHub: {github}
- Vercel: {vercel}
- YouTube: {youtube}
- LinkedIn: {linkedin}
"""
    (folder / "CHECKLIST.md").write_text(checklist, encoding="utf-8")


def write_readme(config: dict, metrics: dict) -> None:
    test = metrics["test"]
    counts = metrics["counts"]
    text = f"""# Waste Image Classification

Deep Learning semester project. MobileNetV2 is fine-tuned on [TrashNet](https://github.com/garythung/trashnet) to classify cardboard, glass, metal, paper, plastic, and trash.

**Test accuracy:** {pct(test['accuracy'])}  
**Macro F1:** {pct(test['macro_f1'])}  
**Images:** {counts['total']} ({counts['train']['total']} train / {counts['val']['total']} validation / {counts['test']['total']} test)

## Load the weights

The checkpoint is a normal Git file at `weights/model.pth` (no Git LFS). It contains `state_dict`, `classes`, `input_size` 224, ImageNet `mean` and `std`, and `arch`.

```python
import torch
from model import build_model

checkpoint = torch.load("weights/model.pth", map_location="cpu", weights_only=False)
model = build_model(num_classes=len(checkpoint["classes"]), pretrained=False)
model.load_state_dict(checkpoint["state_dict"])
model.eval()
```

Class order: cardboard, glass, metal, paper, plastic, trash. The network returns logits.

```bash
python predict.py --image samples/paper.jpg
```

`model_manifest.json` lists the same paths. `weights/model.onnx` is the browser copy. `weights/config.json` records preprocessing: resize the shorter side to 256, center-crop 224, ImageNet normalize.

## Project layout

| Path | Role |
| --- | --- |
| `train.py` | Stratified split, two-phase fine-tune, metrics, ONNX export |
| `model.py` | MobileNetV2 with a 6-class head |
| `predict.py` | JSON prediction for one image |
| `download_dataset.py` | Downloads the resized TrashNet zip |
| `weights/` | Checkpoint, ONNX, class names, config |
| `outputs/metrics.json` | Accuracy, F1, confusion matrix, training history |
| `report/PROJECT_REPORT.md` | Written report |
| `presentation/presentation.pptx` | Slides with speaker notes |
| `presentation/VOICEOVER_SCRIPT.md` | Text to read for the YouTube recording |
| `web/` | Static site deployed on Vercel |

## Reproduce

```bash
python -m venv .venv
.venv\\Scripts\\python -m pip install torch torchvision --index-url https://download.pytorch.org/whl/cu128
.venv\\Scripts\\python -m pip install -r requirements.txt
.venv\\Scripts\\python download_dataset.py
.venv\\Scripts\\python train.py
.venv\\Scripts\\python build_docs.py
node prepare_site.js
```

The CUDA 12.8 wheel is required for an RTX 50-series GPU. `train.py` uses the CPU if CUDA cannot run a kernel. The image archive stays in `data/` and is gitignored.

## Site

`node prepare_site.js` copies the weights, plots, samples, report, and slides into `web/`. Deploy the repository on Vercel. The build command is already in `vercel.json`.

Local preview, after prepare:

```bash
npx --yes serve web -p 4173
```

Then open the demo, `report.html`, and `downloads/model.pth`.

## Dataset citation

Gary Thung and Mindy Yang. TrashNet. Stanford CS 229, 2016. https://github.com/garythung/trashnet

{config['studentName']} · {config['rollNumber']} · {config['course']}
"""
    (ROOT / "README.md").write_text(text, encoding="utf-8")


def main() -> None:
    config = load_json(ROOT / "project.config.json")
    metrics = load_json(ROOT / "outputs" / "metrics.json")
    for name in metrics["class_order"]:
        sample = ROOT / "samples" / f"{name}.jpg"
        if not sample.exists():
            raise SystemExit(f"Missing {sample}. Run train.py first.")

    counts = metrics["counts"]
    test = metrics["test"]
    training = metrics["training"]
    intro = student_intro(config)
    notes = {
        "Title": (
            f"{intro} The title is waste image classification. Say that the network is MobileNetV2, "
            "fine-tuned on the public TrashNet dataset, and that the rest of the talk covers the data, "
            "the training setup, the test numbers, and where the weights live."
        ),
        "Problem": (
            "Explain the problem in plain language. Recycling depends on separating cardboard, paper, "
            "plastic, glass, and metal. The model takes one photograph and picks one of those classes, "
            "or trash when the object does not fit. Pause, then say clearly that this is a course demo "
            "with published weights, not a system that should run a sorting plant."
        ),
        "Objectives": (
            "Walk through the four deliverables. First, TrashNet with a fixed split. Second, the "
            "MobileNetV2 checkpoint and the ONNX file made from it. Third, predict.py, so a test script "
            "can score one image. Fourth, the Vercel page, which shows the demo, the report, and this presentation."
        ),
        "Dataset": (
            f"Credit Gary Thung and Mindy Yang. Say the resized set has {counts['total']} photographs, "
            f"with trash the smallest class at {counts['all']['trash']} images. Give the split: "
            f"{counts['train']['total']} train, {counts['val']['total']} validation, {counts['test']['total']} test, "
            f"seed {metrics['seed']}. Mention the plain posterboard background, because that fact explains "
            "both the accuracy and the limitation you will come back to."
        ),
        "Examples": (
            "Point across the six photos: cardboard, glass, metal, paper, plastic, and trash. "
            "Say these are held-out test images the trained model classified correctly, kept in the repository "
            "so the website has something to click. Mention that every photo is a single object on a light background."
        ),
        "Architecture": (
            f"Name MobileNetV2 and the ImageNet initialization. In one sentence, describe an inverted residual: "
            f"expand, depthwise filter, project back down. Say the old 1,000-way layer was replaced by a 6-way layer, "
            f"and the network has {training['parameter_count']:,} parameters. Close with the input recipe: "
            "shorter side to 256, center crop 224, ImageNet normalization, logits out."
        ),
        "Training": (
            f"Say training used {training['gpu'] or 'the CPU'} and Adam. Mention class weights, because trash "
            f"would otherwise be ignored. Describe the two phases: classifier head first, then the later "
            f"feature blocks. Say {training['epochs_ran']} epochs ran and the best validation accuracy was "
            f"{pct(training['best_val_accuracy'])}. Stress that the test set did not choose the checkpoint."
        ),
        "Results": (
            f"Read the headline slowly: test accuracy {pct(test['accuracy'])}, macro F1 {pct(test['macro_f1'])}. "
            "Explain that macro F1 averages the six classes, so a rare class still counts. Then say "
            + (
                lambda ranked: (
                    f"{ranked[0]} has the highest F1 at {pct(test['per_class'][ranked[0]]['f1'])} "
                    f"on {test['per_class'][ranked[0]]['support']} test images, and {ranked[-1]} has the lowest "
                    f"at {pct(test['per_class'][ranked[-1]]['f1'])} on {test['per_class'][ranked[-1]]['support']} images."
                )
            )(sorted(metrics["class_order"], key=lambda name: test["per_class"][name]["f1"], reverse=True))
        ),
        "Curves": (
            "Describe the plots. The left panel is loss and the right panel is accuracy, for train and validation. "
            "Say the curve flattens, which is why early stopping kept the best validation epoch instead of the last one."
        ),
        "Confusion": (
            f"Tell the audience that rows are the true class and columns are the prediction. Then read this sentence: "
            f"{confusion_sentence(metrics)} Add that glass, plastic, and metal are the mix-ups you would expect "
            "from reflective or transparent objects."
        ),
        "Demo": (
            "This is the slide to slow down on if the evaluation is automated. Say the PyTorch file is "
            "weights/model.pth, the browser file is weights/model.onnx, and model.py rebuilds the architecture. "
            "Read the class order once: cardboard, glass, metal, paper, plastic, trash. Mention predict.py "
            "as the command that prints probabilities for a single image."
        ),
        "Limits": (
            "Be direct. The training photos use a clean background, so a picture of a real bin can fail. "
            f"Trash has only {counts['all']['trash']} images, so that score moves around. End by saying the "
            "demo must not be used to judge hazardous waste."
        ),
        "Close": (
            "Read the dataset credit, then the GitHub, Vercel, YouTube, and LinkedIn links if they are already "
            "filled in. If a link is still a placeholder, say it will be on the submission form. Thank the audience "
            "and stop."
        ),
    }
    write_report(config, metrics)
    write_slides(config, metrics, notes)
    write_voiceover(notes)
    write_submission(config, metrics)
    write_readme(config, metrics)
    print("Wrote report, slides, voiceover, LinkedIn draft, checklist, and README.")


if __name__ == "__main__":
    main()
